"""Configuration contract for the optional hydrogen case study."""

from __future__ import annotations

import ast
import hashlib
import io

import pytest

pytest.importorskip("yaml")

from case_studies.hydrogen import fetch_open_data
from case_studies.hydrogen.workflow import case_config
from case_studies.hydrogen.workflow.foreground_import import (
    normalize_foreground_database,
)
from case_studies.hydrogen.workflow.study_inputs import load_study_inputs


def test_validation_factor_neutral_comparison_preserves_pairs_and_default_values():
    import numpy as np
    import pandas as pd

    module = case_config.CASE_DIR / "workflow/benchmark_comparison.py"
    source = module.read_text(encoding="utf-8")
    helpers = [
        node
        for node in ast.parse(source).body
        if isinstance(node, ast.FunctionDef)
        and node.name in {"factor_neutral_comparison", "comparison_summary"}
    ]
    namespace = {"np": np}
    exec(compile(ast.Module(body=helpers, type_ignores=[]), str(module), "exec"), namespace)
    factor = 1.3108611656722589
    default = pd.DataFrame(
        {
            "dataset_id": [1, 1],
            "indicator": ["example", "example"],
            "type": ["risk", "opportunity"],
            "sense_value": [factor * 2, factor * 20],
            "soca_value": [2.0, 2.0],
            "log_diff": [np.log10(factor), np.log10(factor) + 1],
        }
    )
    original = default.copy(deep=True)
    neutral = namespace["factor_neutral_comparison"](default, factor)
    pd.testing.assert_frame_equal(default, original)
    pd.testing.assert_frame_equal(neutral.iloc[:, :3], default.iloc[:, :3])
    assert neutral["soca_value"].tolist() == [2.0, 2.0]
    assert neutral["sense_value"].tolist() == pytest.approx([2.0, 20.0])
    assert neutral["log_diff"].tolist() == pytest.approx([0.0, 1.0])
    assert namespace["comparison_summary"](neutral)["typical_absolute_factor"] == pytest.approx(
        np.sqrt(10)
    )
    pd.testing.assert_frame_equal(namespace["factor_neutral_comparison"](default, 1.0), default)
    for invalid in [0.0, -1.0, np.nan, np.inf]:
        with pytest.raises(ValueError, match="finite and positive"):
            namespace["factor_neutral_comparison"](default, invalid)


def test_hydrogen_template_resolves_paths_and_credentials(monkeypatch):
    template = case_config.CASE_DIR / "config.yaml.template"
    cfg = case_config.load(template)

    required_paths = {
        "data_dir",
        "ecoinvent_geographies_excel",
        "psilca_path",
        "hydrogen_inventory",
        "methanol_inventory",
        "soca_dir",
    }
    assert required_paths <= cfg["paths"].keys()
    assert "trade_database" not in cfg["paths"]
    assert case_config.path("data_dir", cfg) == case_config.CASE_DIR / "Data"
    assert case_config.TRADE_DATABASE_PATH == (
        case_config.CASE_DIR / "Data" / "regioinvent" / "trade_data.db"
    )
    assert "validation_ecoinvent" not in cfg["versions"]
    assert "validation_ecoinvent_dir" not in cfg["paths"]
    assert "validation_geographies_excel" not in cfg["paths"]
    assert case_config.ecoinvent_database_name(cfg) == "ecoinvent-3.12-cutoff"
    assert (
        case_config.ecoinvent_database_name(cfg, regionalized=True)
        == "ecoinvent-3.12-cutoff - regionalized"
    )

    monkeypatch.setenv("ECOINVENT_USERNAME", "environment-user")
    monkeypatch.setenv("ECOINVENT_PASSWORD", "environment-password")
    assert case_config.ecoinvent_credentials(cfg) == (
        "environment-user",
        "environment-password",
    )

    cfg["ecoinvent_credentials"] = {
        "username": "config-user",
        "password": "config-password",
    }
    assert case_config.require_ecoinvent_credentials(cfg) == (
        "config-user",
        "config-password",
    )


def test_trade_database_fetch_is_verified_and_idempotent(monkeypatch, tmp_path):
    payload = b"test trade database"
    calls = 0

    def open_response(*_args, **_kwargs):
        nonlocal calls
        calls += 1
        return io.BytesIO(payload)

    monkeypatch.setattr(fetch_open_data.urllib.request, "urlopen", open_response)
    monkeypatch.setattr(
        fetch_open_data,
        "TRADE_DATABASE_SHA256",
        hashlib.sha256(payload).hexdigest(),
    )
    destination = tmp_path / "regioinvent" / "trade_data.db"

    fetch_open_data.fetch_trade_database(destination)
    fetch_open_data.fetch_trade_database(destination)

    assert destination.read_bytes() == payload
    assert calls == 1


def test_legacy_foreground_links_are_normalized_without_changing_inventory_values():
    datasets = [
        {
            "database": "legacy-foreground",
            "name": "final product",
            "reference product": "final product",
            "unit": "kg",
            "location": "GLO",
            "exchanges": [
                {
                    "database": "legacy-foreground",
                    "name": "final product",
                    "reference product": "final product",
                    "unit": "kg",
                    "location": "GLO",
                    "type": "production",
                    "amount": 1.0,
                },
                {
                    "database": "legacy-foreground",
                    "name": "internal input",
                    "reference product": "internal input",
                    "unit": "kg",
                    "location": "GLO",
                    "type": "technosphere",
                    "amount": 2.0,
                },
                {
                    "database": "legacy-scenario",
                    "code": "stale-code",
                    "input": ("legacy-scenario", "stale-code"),
                    "name": "external input",
                    "reference product": "external input",
                    "unit": "kg",
                    "location": "RoW",
                    "type": "technosphere",
                    "amount": 3.0,
                },
            ],
        },
        {
            "database": "legacy-foreground",
            "name": "internal input",
            "reference product": "internal input",
            "unit": "kg",
            "location": "GLO",
            "exchanges": [],
        },
    ]

    summary = normalize_foreground_database(datasets, "current-foreground")

    exchanges = datasets[0]["exchanges"]
    assert {dataset["database"] for dataset in datasets} == {"current-foreground"}
    assert exchanges[0]["database"] == "current-foreground"
    assert exchanges[1]["database"] == "current-foreground"
    assert not {"database", "code", "input"} & exchanges[2].keys()
    assert [exchange["amount"] for exchange in exchanges] == [1.0, 2.0, 3.0]
    assert summary.source_databases == ("legacy-foreground",)
    assert (summary.internal_exchanges, summary.external_exchanges) == (2, 1)


def test_foreground_normalization_rejects_duplicate_activity_identities():
    activity = {
        "database": "legacy-foreground",
        "name": "duplicate",
        "reference product": "product",
        "unit": "kg",
        "location": "GLO",
        "exchanges": [],
    }

    with pytest.raises(ValueError, match="duplicate activity identities"):
        normalize_foreground_database([activity.copy(), activity.copy()], "target")


def test_preflight_optional_solver_without_parent_namespace(monkeypatch):
    import runpy

    monkeypatch.syspath_prepend(str(case_config.CASE_DIR))
    helpers = runpy.run_path(str(case_config.CASE_DIR / "check_setup.py"))
    # A dotted lookup raises when its parent is absent, rather than returning None.
    assert helpers["_module_available"]("sense_nonexistent_solver_namespace.umfpack") is False
    assert helpers["_module_available"]("json") is True


def _synthetic_rules():
    return {
        "schema": 1,
        "ecoinvent_version": "test",
        "subset_targets": [["new", "p", "AA"]],
        "activity_specs": [["new", "p", 1]],
        "methanol_replacement": {
            "target_activity_name": "example",
            "target_reference_product": "p",
            "old_input_names": ["old"],
            "old_input_reference_product": "p",
            "new_input_name": "new",
            "new_input_reference_product": "p",
        },
        "co2_flow": {"name": "example flow", "categories": ["example"], "keyword": "example"},
    }


def _study_workbook(path):
    import openpyxl

    rules = _synthetic_rules()
    wb = openpyxl.Workbook()
    inventory = wb.active
    inventory.title = "Inventory"
    for row in [
        ["Database", "synthetic"],
        [],
        ["Activity", "Example"],
        ["reference product", "Product"],
        ["unit", "kg"],
        [],
        ["Exchanges"],
        ["name", "amount", "type"],
        ["Example", 1, "production"],
    ]:
        inventory.append(row)
    activities = wb.create_sheet("SENSE activities")
    activities.append(["skip", "Study selections"])
    activities.append(["purpose", "name", "reference product", "location", "demand"])
    activities.append(["subset", "new", "p", "AA", None])
    activities.append(["functional unit", "new", "p", None, 1])
    settings = wb.create_sheet("SENSE settings")
    settings.append(["skip", "Study settings"])
    settings.append(["setting", "value"])
    settings.append(["schema", 1])
    settings.append(["ecoinvent_version", "test"])
    for prefix, values in [
        ("methanol.", rules["methanol_replacement"]),
        ("co2.", rules["co2_flow"]),
    ]:
        for key, value in values.items():
            for item in value if isinstance(value, list) else [value]:
                settings.append([prefix + key, item])
    wb.save(path)
    return wb


def test_workbook_study_selections_roundtrip_and_importer_separation(tmp_path):
    from bw2io.extractors.excel import ExcelExtractor
    from bw2io.importers.excel import ExcelImporter, valid_first_cell

    path = tmp_path / "inventory.xlsx"
    with pytest.raises(FileNotFoundError, match="workbook"):
        load_study_inputs(path, ecoinvent_version="test")
    _study_workbook(path).close()
    assert load_study_inputs(path, ecoinvent_version="test") == _synthetic_rules()
    extracted = [
        (name, rows) for name, rows in ExcelExtractor.extract(path) if valid_first_cell(name, rows)
    ]
    assert [name for name, _ in extracted] == ["Inventory"]
    importer = ExcelImporter.__new__(ExcelImporter)
    importer.db_name = "synthetic"
    records = importer.process_activities(extracted)
    assert len(records) == 1 and records[0]["name"] == "Example"
    assert records[0]["exchanges"][0]["amount"] == 1
    with pytest.raises(ValueError, match="version"):
        load_study_inputs(path, ecoinvent_version="other")


@pytest.mark.parametrize(
    "invalid",
    ["missing_sheet", "missing_skip", "duplicate_setting", "zero_demand", "unknown_purpose"],
)
def test_invalid_workbook_metadata_stops_before_the_study(tmp_path, invalid):
    path = tmp_path / "inventory.xlsx"
    wb = _study_workbook(path)
    if invalid == "missing_sheet":
        del wb["SENSE settings"]
    elif invalid == "missing_skip":
        wb["SENSE activities"]["A1"] = "Study selections"
    elif invalid == "duplicate_setting":
        wb["SENSE settings"].append(["schema", 1])
    elif invalid == "zero_demand":
        wb["SENSE activities"]["E4"] = 0
    else:
        wb["SENSE activities"]["A4"] = "unknown"
    wb.save(path)
    wb.close()
    with pytest.raises(ValueError):
        load_study_inputs(path, ecoinvent_version="test")


@pytest.mark.parametrize("reuse", [False, True])
def test_preflight_rejects_unprepared_inventory_before_reporting_success(
    tmp_path, monkeypatch, capsys, reuse
):
    from case_studies.hydrogen import check_setup

    path = tmp_path / "inventory.xlsx"
    workbook = _study_workbook(path)
    del workbook["SENSE settings"]
    workbook.save(path)
    workbook.close()
    cfg = {"versions": {"ecoinvent": "test"}, "paths": {"hydrogen_inventory": str(path)}}
    monkeypatch.setattr(check_setup, "load", lambda: cfg)
    monkeypatch.setattr(check_setup, "_check_environment", lambda *a, **kw: True)
    monkeypatch.setattr(check_setup, "_main_inputs", lambda *a, **kw: ())
    args = ["--reuse-backgrounds"] if reuse else []
    assert check_setup.main(args) == 1
    output = capsys.readouterr().out
    assert "missing 'SENSE settings'" in output
    assert "Setup is complete." not in output

    _study_workbook(path).close()
    assert check_setup.main(args) == 0
    assert "1 functional-unit selections" in capsys.readouterr().out


def test_preflight_rejects_wrong_workbook_version(tmp_path, capsys):
    from case_studies.hydrogen.check_setup import _check_study_workbook

    path = tmp_path / "inventory.xlsx"
    _study_workbook(path).close()
    cfg = {"versions": {"ecoinvent": "other"}, "paths": {"hydrogen_inventory": str(path)}}
    assert not _check_study_workbook(cfg)
    assert "does not match" in capsys.readouterr().out


def test_manuscript_versions_are_fixed(tmp_path):
    import yaml

    template = case_config.CASE_DIR / "config.yaml.template"
    config = yaml.safe_load(template.read_text(encoding="utf-8"))
    config["versions"]["psilca"] = "3.0"
    local = tmp_path / "config.local.yaml"
    local.write_text(yaml.safe_dump(config), encoding="utf-8")
    with pytest.raises(ValueError, match="manuscript case study uses"):
        case_config.load(local)
