"""Read private study selections from the supplied hydrogen inventory workbook."""

from __future__ import annotations

import math
from pathlib import Path
from typing import Any

import openpyxl


def load_study_inputs(workbook: Path, *, ecoinvent_version: str) -> dict[str, Any]:
    """Validate the two study sheets without importing inventories into Brightway.

    The source workbook contains version-correct inventory links. Its additional
    sheets identify the original study's selections; they do not define runtime
    migrations. Both sheets start with ``skip`` so Brightway's Excel importer
    ignores them. Exact database identities remain private input data.
    """
    workbook = Path(workbook)
    if not workbook.is_file():
        raise FileNotFoundError(f"Hydrogen inventory workbook not found: {workbook}")
    wb = openpyxl.load_workbook(workbook, read_only=True, data_only=True)
    try:

        def table(sheet: str, headers: tuple[str, ...]) -> list[tuple[Any, ...]]:
            if sheet not in wb.sheetnames:
                raise ValueError(
                    f"{workbook.name} is missing '{sheet}'. Use the prepared inventory "
                    "workbook for this study; see the case-study README."
                )
            ws = wb[sheet]
            rows = ws.iter_rows(values_only=True)
            first = next(rows, ())
            if not first or str(first[0]).strip().lower() != "skip":
                raise ValueError(f"'{sheet}' must have 'skip' in A1 for the Brightway importer.")
            heading = next(rows, ())
            if tuple(heading[: len(headers)]) != headers or any(heading[len(headers) :]):
                raise ValueError(f"Unexpected columns in '{sheet}'; expected {headers}.")
            data = []
            for row in rows:
                if any(value is not None for value in row):
                    if any(value is not None for value in row[len(headers) :]):
                        raise ValueError(f"Unexpected extra values in '{sheet}'.")
                    data.append(tuple(row[: len(headers)]))
            return data

        settings_rows = table("SENSE settings", ("setting", "value"))
        activity_rows = table(
            "SENSE activities", ("purpose", "name", "reference product", "location", "demand")
        )
    finally:
        wb.close()

    settings: dict[str, list[Any]] = {}
    for key, value in settings_rows:
        if not isinstance(key, str) or not key.strip():
            raise ValueError("Study setting names must be nonempty strings.")
        settings.setdefault(key, []).append(value)
    methanol_fields = {
        "target_activity_name",
        "target_reference_product",
        "old_input_names",
        "old_input_reference_product",
        "new_input_name",
        "new_input_reference_product",
    }
    required = {"schema", "ecoinvent_version", "co2.name", "co2.keyword", "co2.categories"}
    required.update("methanol." + key for key in methanol_fields)
    if set(settings) != required:
        raise ValueError("Study settings must contain exactly the fields documented in the README.")
    repeated = {"methanol.old_input_names", "co2.categories"}
    if any(len(values) != 1 for key, values in settings.items() if key not in repeated):
        raise ValueError("Only old-input names and CO2 categories may use repeated setting rows.")
    if settings["schema"] != [1] or isinstance(settings["schema"][0], bool):
        raise ValueError("Unsupported study input schema; expected 1.")
    if settings["ecoinvent_version"] != [str(ecoinvent_version)]:
        raise ValueError("Study workbook does not match the configured ecoinvent version.")

    def text(value: Any) -> bool:
        return isinstance(value, str) and bool(value.strip())

    if any(
        not text(value) for key, values in settings.items() if key != "schema" for value in values
    ):
        raise ValueError("Study setting values other than schema must be nonempty text.")
    subsets, specs = [], []
    for purpose, name, product, location, demand in activity_rows:
        if not text(name) or not text(product):
            raise ValueError(
                "Each study activity must specify its exact name and reference product."
            )
        if purpose == "subset" and text(location) and demand is None:
            subsets.append([name, product, location])
        elif (
            purpose == "functional unit"
            and location is None
            and not isinstance(demand, bool)
            and isinstance(demand, (int, float))
            and math.isfinite(demand)
            and demand > 0
        ):
            specs.append([name, product, demand])
        else:
            raise ValueError(
                "Use 'subset' with a location, or 'functional unit' with a positive finite demand."
            )
    if not subsets or not specs:
        raise ValueError(
            "Study activities must include both subset and functional-unit selections."
        )
    if len(set(map(tuple, subsets))) != len(subsets) or len(
        {tuple(row[:2]) for row in specs}
    ) != len(specs):
        raise ValueError("Duplicate study activity selections are not allowed.")
    return {
        "schema": 1,
        "ecoinvent_version": str(ecoinvent_version),
        "subset_targets": subsets,
        "activity_specs": specs,
        "methanol_replacement": {
            key: (
                settings["methanol." + key]
                if key == "old_input_names"
                else settings["methanol." + key][0]
            )
            for key in methanol_fields
        },
        "co2_flow": {
            "name": settings["co2.name"][0],
            "keyword": settings["co2.keyword"][0],
            "categories": settings["co2.categories"],
        },
    }
