"""Write aggregate SOCA comparison coverage without publishing inventory lists."""

import json
from pathlib import Path
from zipfile import ZipFile

import pandas as pd
from openpyxl import load_workbook

from . import case_config
from .si_common import require, write_text
from .study_protocol import load_protocol


def write_comparison_coverage(tables, data_dir: Path, output_dir: Path) -> dict:
    """Summarize positive comparison pairs and one-sided exchange coverage."""
    universe = set(
        pd.read_parquet(data_dir / "cache/soca/mapping_df.parquet", columns=["dataset_id"])[
            "dataset_id"
        ]
        .dropna()
        .astype("int64")
    )
    soca_ids = set(
        pd.read_parquet(data_dir / "cache/soca/xml_df.parquet", columns=["dataset_id"])[
            "dataset_id"
        ]
        .dropna()
        .astype("int64")
    )
    require(soca_ids <= universe, "SOCA coverage lies outside the exported activity universe")
    counts = {}
    for mean in ("geometric", "arithmetic"):
        table = tables[mean]
        book = load_workbook(
            data_dir / f"sense_soca_activity_coverage_{mean}.xlsx",
            read_only=True,
            data_only=True,
        )
        try:
            exclusive = {}
            for name in ("Only_SENSE", "Only_SOCA"):
                rows = book[name].iter_rows(values_only=True)
                header = next(rows)
                index = header.index("dataset_id")
                exclusive[name] = {int(row[index]) for row in rows if row[index] is not None}
            require(exclusive["Only_SOCA"] <= soca_ids, "Invalid SOCA-only coverage")
            require(
                not (exclusive["Only_SENSE"] & soca_ids),
                "SENSE-only activities overlap SOCA coverage",
            )
            require(
                exclusive["Only_SENSE"] <= universe,
                "SENSE-only coverage outside the activity universe",
            )
            both = soca_ids - exclusive["Only_SOCA"]
            require(
                set(table["dataset_id"].astype("int64")) <= both,
                "Compared activities lack shared coverage",
            )
            counts[mean] = {
                "universe": len(universe),
                "both_exchange_coverage": len(both),
                "positive_pairs": len(table),
                "activities_with_positive_pairs": int(table["dataset_id"].nunique()),
                **{name: len(ids) for name, ids in exclusive.items()},
                "neither_exchange_coverage": len(universe - (soca_ids | exclusive["Only_SENSE"])),
            }
        finally:
            book.close()

    rows = [
        ("SOCA export/activity-overview universe", "universe"),
        ("Activities with exchanges in both implementations", "both_exchange_coverage"),
        ("Positive matched activity--indicator pairs", "positive_pairs"),
        ("Activities with at least one positive matched pair", "activities_with_positive_pairs"),
        ("Activities with SENSE exchanges only", "Only_SENSE"),
        ("Activities with SOCA exchanges only", "Only_SOCA"),
        ("Activities with no extracted social exchanges in either", "neither_exchange_coverage"),
    ]
    common_counts = counts["geometric"] == counts["arithmetic"]
    count_note = " Counts are identical for both aggregation variants." if common_counts else ""
    lines = [
        r"\begin{table}[tbp]",
        r"\centering",
        r"\caption{Numerical comparison population and one-sided social-exchange coverage for the SENSE--SOCA benchmark. Positive comparisons require compatible activity and indicator identities and strictly positive intensities in both implementations. One-sided coverage means at least one extracted social exchange and is not a count of positive numerical comparisons."
        + count_note
        + "}",
        r"\label{tab:comparison-coverage}",
        r"\begin{tabular}{lr}" if common_counts else r"\begin{tabular}{lrr}",
        r"\toprule",
        r"Population & Count \\" if common_counts else r"Population & Geometric & Arithmetic \\",
        r"\midrule",
    ]
    for label, key in rows:
        values = [f"{counts['geometric'][key]:,}"]
        if not common_counts:
            values.append(f"{counts['arithmetic'][key]:,}")
        lines.append(" & ".join([label, *values]) + r" \\")
    lines.extend([r"\bottomrule", r"\end{tabular}", r"\end{table}", ""])
    write_text("\n".join(lines), output_dir / "comparison_coverage.tex")
    return counts


def run():
    cfg = case_config.load()
    data = case_config.path("data_dir", cfg)
    spec = load_protocol()
    context = {
        "versions": cfg["versions"],
        "projects": cfg["brightway"],
        "social_layers": spec["social_layers"],
        "sense": spec["sense"],
    }
    snapshot = data / "figure_inputs/validation.zip"
    # This table reports the population of a saved benchmark, not a new run of
    # today's analysis code. Preserve its original provenance when reading it.
    with ZipFile(snapshot) as archive:
        manifest = json.loads(archive.read("manifest.json"))
        require(
            manifest.get("schema") == 1 and manifest["metadata"].get("context") == context,
            "Benchmark snapshot belongs to another study or has an unsupported format",
        )
        tables = {
            mean: pd.read_csv(archive.open(mean + ".csv"), usecols=["dataset_id"])
            for mean in ("geometric", "arithmetic")
        }
    counts = write_comparison_coverage(tables, data, data / "si")
    from sense.artifacts import write_json_atomic
    from sense.utils.cache import file_hash

    sources = [
        snapshot,
        data / "cache/soca/mapping_df.parquet",
        data / "cache/soca/xml_df.parquet",
        *[
            data / f"sense_soca_activity_coverage_{mean}.xlsx"
            for mean in ("geometric", "arithmetic")
        ],
    ]
    write_json_atomic(
        data / "si/comparison_coverage.json",
        {
            "counts": counts,
            "benchmark_created_utc": manifest["created_utc"],
            "benchmark_analysis_signature": manifest["analysis_signature"],
            "input_sha256": {str(path.relative_to(data)): file_hash(path) for path in sources},
        },
    )
    print("Wrote SOCA comparison coverage table from saved benchmark data.")
